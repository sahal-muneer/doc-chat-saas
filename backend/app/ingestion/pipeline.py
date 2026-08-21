"""
pipeline.py
===========

WHY THIS FILE EXISTS:

parser.py, chunker.py, and embedder.py are all "pure" in a specific sense:
none of them know the others exist. parser.py's job ends the moment it
hands back a list of (page_number, text) tuples. chunker.py's job is to
turn ONE string of text into Chunks — it has no concept of "a document
with multiple pages" at all. embedder.py's job is to turn Chunks that
already exist into vectors — it has no idea how those Chunks were produced.
Nothing in the codebase actually calls all three together. This file is
that missing piece: given a file on disk, run the WHOLE "extract → chunk →
embed" pipeline and hand back one flat list of fully embedded Chunks,
ready to be written into pgvector (storage isn't built yet — that's next).

WHY EMBEDDING HAPPENS HERE, AS THE LAST STEP, RATHER THAN CHUNK-BY-CHUNK
INSIDE THE PAGE LOOP BELOW: embedder.py's own docstring explains that
embedding many texts in ONE batched call is dramatically more efficient
than embedding them one at a time. If we called embed_chunks() once per
page inside the loop, we'd be doing several small, less-efficient batches
(one per page) instead of one large, efficient batch covering the whole
document. So this file waits until every page has been parsed and chunked,
then embeds everything in a single call at the very end.

THE ONE REAL PROBLEM THIS FILE HAS TO SOLVE: chunk numbering across pages.

chunk_text() numbers the chunks IT creates starting from `start_index`
(see the docstring on that parameter in chunker.py for the full reasoning).
But chunk_text() only ever sees one page's text at a time — it has no idea
how many chunks earlier pages already produced. So THIS file is what tracks
that running total: after chunking page 1, we know exactly how many chunks
it produced, and we tell chunk_text() to start page 2's numbering right
after that. Get this wrong (e.g. just call chunk_text() once per page with
no start_index) and every page's first chunk comes out numbered 0 — a
document with 5 pages would end up with FIVE different chunks all claiming
to be "chunk 0," which would silently corrupt anything downstream that
assumes chunk_index is unique within a document (e.g. using it to order
chunks back into original document order, or as part of a database primary
key).
"""

from pathlib import Path

from app.ingestion.chunker import Chunk, chunk_text
from app.ingestion.embedder import embed_chunks
from app.ingestion.parser import parse_document


def ingest_document(
    file_path: str | Path,
    document_id: str,
    chunk_size_tokens: int = 600,
    overlap_tokens: int = 80,
) -> list[Chunk]:
    """
    Run the full extract → chunk → embed pipeline for one document.

    Parses the file into (page_number, text) pairs, chunks each page's text
    individually (so every Chunk keeps an accurate page_number for source
    citations — see the parser.py and chunker.py docstrings for why that
    matters), numbers every chunk continuously from 0 across the WHOLE
    document rather than restarting per page, then embeds all of them in
    one batched call. Every Chunk that comes back from this function has
    its `.embedding` field filled in — it's fully ready to be written into
    the vector database, no further processing needed.

    HEAVY OPERATION WARNING, worth knowing before calling this casually:
    the first call to this function in a running process will trigger
    embedder.py's lazy model load — potentially several seconds to load
    bge-m3's ~2GB of weights into memory. Every call after that, in the
    same running process, reuses the already-loaded model and is much
    faster. This is exactly the "load once, reuse many times" pattern from
    embedder.py's docstring, just experienced here from the caller's side.

    chunk_size_tokens and overlap_tokens are exposed here — instead of
    hidden behind a generic **kwargs passthrough to chunk_text() — as an
    explicit, deliberate choice: it means these two lines duplicate two
    parameter names that also exist in chunker.py, which is a little more
    to keep in sync if chunk_text() ever grows new tuning knobs. The
    trade-off is worth it here: this is a learning codebase, and named
    parameters mean anyone calling ingest_document() gets autocomplete,
    type checking, and can *see* what's tunable without having to go read
    chunk_text()'s signature first. **kwargs would technically be less
    duplication, at the cost of that visibility.
    """
    pages = parse_document(file_path)

    all_chunks: list[Chunk] = []
    next_index = 0

    for page_number, page_text in pages:
        # start_index=next_index is the fix described in the module
        # docstring above: tell this call exactly where to continue
        # numbering from, based on how many chunks every PRIOR page in
        # this same document has already produced.
        page_chunks = chunk_text(
            page_text,
            document_id=document_id,
            chunk_size_tokens=chunk_size_tokens,
            overlap_tokens=overlap_tokens,
            page_number=page_number,
            start_index=next_index,
        )
        all_chunks.extend(page_chunks)

        # Advance by however many chunks THIS page actually produced — not
        # a fixed guess. A short page might produce 1 chunk; a long one
        # might produce 5. next_index has to reflect reality, not an
        # assumption about page size.
        next_index += len(page_chunks)

    # One batched call for the whole document — see the module docstring
    # for why this happens out here, after the loop, rather than once per
    # page inside it.
    return embed_chunks(all_chunks)


if __name__ == "__main__":
    # Manual test: build a small demo .docx with two pages, each containing
    # enough text to produce MULTIPLE chunks per page (using a small
    # chunk_size_tokens). This is deliberately different from parser.py's
    # own demo, which used tiny one-sentence pages — those would only ever
    # produce exactly one chunk per page, which wouldn't actually prove
    # anything about the chunk_index-continuity fix this file exists for.
    # Here, if the fix works, chunk_index should climb 0, 1, 2, 3, ...
    # smoothly across the page boundary with no reset back to 0.
    import tempfile

    from docx import Document as DocxDocument

    tmp_dir = Path(tempfile.gettempdir())
    docx_path = tmp_dir / "pipeline_demo.docx"

    demo_doc = DocxDocument()
    demo_doc.add_paragraph(
        "Retrieval-Augmented Generation combines information retrieval with "
        "text generation. It retrieves relevant text from a knowledge base "
        "before generating an answer. This grounds the model's response in "
        "real, verifiable, up to date information rather than only what it "
        "memorized during training. RAG is especially useful for private "
        "documents the model has never seen."
    )
    demo_doc.add_page_break()
    demo_doc.add_paragraph(
        "Vector databases store embeddings so that semantically similar "
        "pieces of text can be found quickly. Postgres with the pgvector "
        "extension is one popular, simple choice for a first version. "
        "Qdrant is a dedicated vector database with more advanced indexing "
        "options built for very large datasets and higher query volume."
    )
    demo_doc.save(docx_path)

    chunks = ingest_document(
        docx_path,
        document_id="demo-doc",
        chunk_size_tokens=20,
        overlap_tokens=5,
    )

    for chunk in chunks:
        print(
            f"chunk_index={chunk.chunk_index}  page={chunk.page_number}  "
            f"tokens={chunk.token_count}  embedding_dims={len(chunk.embedding)}"
        )
        print(f"  {chunk.text!r}")
        print()
